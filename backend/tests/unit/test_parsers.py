"""
Parser tests. Test files are generated on the fly (PyMuPDF can write PDFs,
python-docx can write DOCX), so the tests don't depend on sample_data/.
"""

from pathlib import Path

import docx
import pymupdf
import pytest

from app.services.ingestion.parsers import ParseError, parse_file


def make_pdf(path: Path, pages: list[list[tuple[str, float]]], title: str = "") -> Path:
    """Each page is a list of (text, font_size) lines, drawn top to bottom."""
    doc = pymupdf.open()
    for lines in pages:
        page = doc.new_page()
        y = 72
        for text, size in lines:
            page.insert_text((72, y), text, fontsize=size, fontname="helv")
            y += size * 2.2
    if title:
        doc.set_metadata({"title": title})
    doc.save(path)
    return path


def test_pdf_headings_pages_and_title(tmp_path):
    pdf = make_pdf(
        tmp_path / "lease.pdf",
        [
            [("Lease Agreement", 18), ("This lease is between the parties below.", 11)],
            [("Termination", 14), ("The tenant must give 60 days written notice.", 11)],
        ],
    )
    parsed = parse_file(pdf, "lease.pdf")
    assert parsed.page_count == 2
    headings = [(b.text, b.level, b.page) for b in parsed.blocks if b.kind == "heading"]
    assert headings == [("Lease Agreement", 1, 1), ("Termination", 2, 2)]
    notice = next(b for b in parsed.blocks if "60 days" in b.text)
    assert notice.page == 2 and notice.kind == "paragraph"
    assert parsed.title == "Lease Agreement"


def test_pdf_metadata_title_preferred_unless_junk(tmp_path):
    body = [[("Unit 1A", 18), ("Inspection text.", 11)]]
    good = make_pdf(tmp_path / "a.pdf", body, title="Inspection Reports 2026")
    junk = make_pdf(tmp_path / "b.pdf", body, title="Microsoft Word - draft3.docx")
    assert parse_file(good, "a.pdf").title == "Inspection Reports 2026"
    assert parse_file(junk, "b.pdf").title == "Unit 1A"


def test_scanned_pdf_without_text_fails_clearly(tmp_path):
    doc = pymupdf.open()
    doc.new_page()
    doc.save(tmp_path / "scan.pdf")
    with pytest.raises(ParseError, match="scanned"):
        parse_file(tmp_path / "scan.pdf", "scan.pdf")


def test_corrupt_pdf_raises_parse_error(tmp_path):
    bad = tmp_path / "bad.pdf"
    bad.write_bytes(b"%PDF-1.7 this is not really a pdf")
    with pytest.raises(ParseError):
        parse_file(bad, "bad.pdf")


def test_docx_heading_styles_and_tables(tmp_path):
    document = docx.Document()
    document.add_heading("Maintenance Policy", level=1)
    document.add_paragraph("Urgent repairs are handled quickly.")
    table = document.add_table(rows=2, cols=2)
    table.rows[0].cells[0].text = "Repair"
    table.rows[0].cells[1].text = "Response time"
    table.rows[1].cells[0].text = "Burst pipe"
    table.rows[1].cells[1].text = "4 hours"
    path = tmp_path / "policy.docx"
    document.save(path)

    parsed = parse_file(path, "policy.docx")
    assert parsed.blocks[0].kind == "heading" and parsed.blocks[0].text == "Maintenance Policy"
    assert parsed.page_count is None
    assert any(b.text == "Repair: Burst pipe; Response time: 4 hours" for b in parsed.blocks)


def test_markdown_headings_ignore_code_fences(tmp_path):
    md = tmp_path / "guide.md"
    md.write_text(
        "# API Guide\n\nIntro text.\n\n## Auth\n\nSend a token.\n\n"
        "```bash\n# not a heading\ncurl ...\n```\n"
    )
    parsed = parse_file(md, "guide.md")
    headings = [(b.text, b.level) for b in parsed.blocks if b.kind == "heading"]
    assert headings == [("API Guide", 1), ("Auth", 2)]
    assert any("# not a heading" in b.text for b in parsed.blocks)


def test_text_file_title_heuristics(tmp_path):
    txt = tmp_path / "guide.txt"
    txt.write_text(
        "TENANT WELCOME GUIDE\n\n1. COLLECTING YOUR KEYS\n\n"
        "Keys are collected from the office. Bring photo ID.\n\n"
        "This short line is a sentence.\n\nMore text follows here."
    )
    parsed = parse_file(txt, "guide.txt")
    kinds = [(b.kind, b.text[:20]) for b in parsed.blocks]
    assert kinds[0] == ("heading", "TENANT WELCOME GUIDE")
    assert kinds[1] == ("heading", "1. COLLECTING YOUR K")
    assert ("paragraph", "This short line is a") in kinds


def test_unsupported_extension(tmp_path):
    exe = tmp_path / "x.exe"
    exe.write_bytes(b"MZ")
    with pytest.raises(ParseError, match="Unsupported"):
        parse_file(exe, "x.exe")


def test_html_parser_reads_headings_tables_and_drops_navigation(tmp_path):
    html = tmp_path / "wiki.html"
    html.write_text(
        "<html><head><title>Support Wiki</title></head><body>"
        "<nav><a href='/'>Home</a><a href='/pricing'>Pricing</a></nav>"
        "<script>track('page');</script>"
        "<h1>Refund Policy</h1><p>Refunds are issued within 5 business days.</p>"
        "<h2>Exceptions</h2><ul><li>Digital goods are non-refundable.</li></ul>"
        "<table><tr><th>Plan</th><th>Fee</th></tr><tr><td>Premium</td><td>$9</td></tr></table>"
        "<footer>&copy; Kestrel Pay</footer></body></html>"
    )
    parsed = parse_file(html, "wiki.html")
    texts = [(b.kind, b.text) for b in parsed.blocks]

    assert parsed.title == "Support Wiki"
    assert ("heading", "Refund Policy") in texts and ("heading", "Exceptions") in texts
    assert [b.level for b in parsed.blocks if b.kind == "heading"] == [1, 2]
    assert ("paragraph", "Plan: Premium; Fee: $9") in texts  # table row with column names
    joined = " ".join(b.text for b in parsed.blocks)
    assert "Pricing" not in joined and "track(" not in joined and "Kestrel Pay" not in joined


def test_html_without_text_fails_clearly(tmp_path):
    empty = tmp_path / "app.html"
    empty.write_text("<html><body><script>render()</script></body></html>")
    with pytest.raises(ParseError, match="No readable text"):
        parse_file(empty, "app.html")


def test_csv_rows_become_readable_lines(tmp_path):
    csv_file = tmp_path / "rent_roll.csv"
    csv_file.write_text("Unit,Monthly rent,Lease start\n4B,2450,2026-03-01\n7A,3100,2025-11-01\n\n")
    parsed = parse_file(csv_file, "rent_roll.csv")
    texts = [b.text for b in parsed.blocks]

    assert parsed.title == "Rent Roll"
    assert texts[0] == "Rent Roll" and parsed.blocks[0].kind == "heading"
    assert "Columns: Unit, Monthly rent, Lease start" in texts
    assert "Unit: 4B; Monthly rent: 2450; Lease start: 2026-03-01" in texts
    assert len(texts) == 4  # heading + columns + 2 rows (the blank line is skipped)


def test_csv_with_semicolons_and_no_data_rows(tmp_path):
    semicolons = tmp_path / "fees.csv"
    semicolons.write_text("Fee;Amount\nCard replacement;10\n")
    assert "Fee: Card replacement; Amount: 10" in [
        b.text for b in parse_file(semicolons, "fees.csv").blocks
    ]

    header_only = tmp_path / "empty.csv"
    header_only.write_text("Fee,Amount\n")
    with pytest.raises(ParseError, match="no data rows"):
        parse_file(header_only, "empty.csv")
