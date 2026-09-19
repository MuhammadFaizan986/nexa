#!/usr/bin/env python3
"""
generate_sample_data.py - build NEXA's synthetic demo documents.

WHY THIS SCRIPT EXISTS
----------------------
NEXA needs realistic documents to ingest, search and evaluate against, but we
must never use real client or personal data in demos (plan, Section 13). So we
write fictional documents by hand as simple Markdown "sources" and let this
script turn them into the formats a real client would upload: PDF, DOCX, MD and
TXT. Having several formats on purpose exercises every parser in the ingestion
pipeline.

WHERE THINGS LIVE
-----------------
    sample_data/_source/<domain>/*.md|*.txt   hand-written sources (edit these)
    sample_data/<domain>/*                    generated outputs (do not edit;
                                              re-run this script instead)

One document is not hand-written: `inspection_reports_2026.pdf` is built from a
seeded random generator (see build_inspection_reports_source) so we get a long,
50-page PDF with a unique set of facts on every page, which is exactly what the
Week 1 "upload a 50-page PDF and check the page numbers" test needs.

HOW TO RUN
----------
It needs Python 3.12 + PyMuPDF + python-docx. The easiest way is Docker:

    docker run --rm -v "$PWD":/work -w /work python:3.12-slim \
        sh -c "pip install -q pymupdf python-docx && python scripts/generate_sample_data.py"

DETERMINISM
-----------
Running the script twice produces byte-identical files. That matters because
the backend deduplicates uploads by SHA-256 hash: if regenerating changed the
bytes (e.g. an embedded "created at" timestamp), re-seeding the demo tenants
would create duplicate documents instead of being skipped.

THE MINI MARKDOWN FORMAT USED BY THE SOURCES
--------------------------------------------
    # Heading 1            -> big bold heading (PDF 18pt / Word "Heading 1")
    ## Heading 2           -> smaller bold heading (PDF 14pt / Word "Heading 2")
    ### Heading 3          -> Word "Heading 3" (DOCX sources only)
    - bullet item          -> a bullet point (one line in the source)
    | a | b |              -> a table row (DOCX sources only)
    <!-- pagebreak -->     -> start a new PDF page (PDF sources only)
    blank line             -> ends a paragraph; consecutive lines are joined
"""

from __future__ import annotations

import datetime as dt
import hashlib
import io
import random
import zipfile
from dataclasses import dataclass, field
from pathlib import Path

import docx  # python-docx
import pymupdf  # PyMuPDF (a.k.a. fitz)

ROOT = Path(__file__).resolve().parents[1]
SOURCE_DIR = ROOT / "sample_data" / "_source"
OUTPUT_DIR = ROOT / "sample_data"

PAGE_BREAK = "<!-- pagebreak -->"
DISCLAIMER = "Synthetic demo document - fictional company, not real data."

# A fixed timestamp written into every file's metadata (see DETERMINISM above).
FIXED_TIMESTAMP = dt.datetime(2026, 1, 1, 0, 0, 0)


# ---------------------------------------------------------------------------
# Which documents exist, in which format, and which running header they get.
# ---------------------------------------------------------------------------


@dataclass(frozen=True)
class DocSpec:
    domain: str  # "property" | "fintech" | "company" -> output sub-folder
    output: str  # output file name; its extension decides the renderer
    title: str  # document title stored in PDF/DOCX metadata
    source: str | None = None  # file in _source/<domain>/ (None = generated)
    # PDFs only: a small line printed at the top of EVERY page. Real company
    # PDFs almost always have one. It exists here so the ingestion cleaner's
    # "repeated header/footer removal" step has something real to remove.
    running_header: str | None = None


HARBOURVIEW_HEADER = "Harbourview Property Group - Confidential"
KESTREL_HEADER = "Kestrel Pay - Customer Information"
LUMEN_HEADER = "Lumen Labs - Internal - Employee Handbook v4.2"

DOCUMENTS: list[DocSpec] = [
    # --- Property (fictional "Harbourview Property Group") -----------------
    DocSpec("property", "lease_unit_4b.pdf", "Residential Tenancy Agreement - Unit 4B",
            source="lease_unit_4b.md", running_header=HARBOURVIEW_HEADER),
    DocSpec("property", "lease_unit_7a.pdf", "Residential Tenancy Agreement - Unit 7A",
            source="lease_unit_7a.md", running_header=HARBOURVIEW_HEADER),
    DocSpec("property", "building_c_rules.md", "Building C House Rules",
            source="building_c_rules.md"),
    DocSpec("property", "maintenance_policy.docx", "Maintenance and Repairs Policy",
            source="maintenance_policy.md"),
    DocSpec("property", "inspection_reports_2026.pdf", "Routine Inspection Reports 2026",
            source=None, running_header=HARBOURVIEW_HEADER),
    DocSpec("property", "tenant_welcome_guide.txt", "Tenant Welcome Guide",
            source="tenant_welcome_guide.txt"),
    # --- Fintech (fictional "Kestrel Pay") ----------------------------------
    DocSpec("fintech", "fee_schedule.pdf", "Kestrel Pay Fee Schedule",
            source="fee_schedule.md", running_header=KESTREL_HEADER),
    DocSpec("fintech", "kyc_aml_policy.docx", "KYC and AML/CTF Policy",
            source="kyc_aml_policy.md"),
    DocSpec("fintech", "payment_error_codes.md", "Payment Error Codes",
            source="payment_error_codes.md"),
    DocSpec("fintech", "customer_faq.md", "Customer FAQ",
            source="customer_faq.md"),
    DocSpec("fintech", "internal_chargeback_procedure.txt", "Disputes and Chargebacks Procedure",
            source="internal_chargeback_procedure.txt"),
    # --- Company (fictional "Lumen Labs") -----------------------------------
    DocSpec("company", "employee_handbook.pdf", "Lumen Labs Employee Handbook",
            source="employee_handbook.md", running_header=LUMEN_HEADER),
    DocSpec("company", "it_security_policy.docx", "IT Security Policy",
            source="it_security_policy.md"),
    DocSpec("company", "api_documentation.md", "Lumen Platform API v2 Reference",
            source="api_documentation.md"),
    DocSpec("company", "release_notes.md", "Lumen Platform Release Notes",
            source="release_notes.md"),
    DocSpec("company", "onboarding_sop.txt", "New Starter Onboarding SOP",
            source="onboarding_sop.txt"),
]


# ---------------------------------------------------------------------------
# A tiny Markdown parser. We only need a handful of constructs, so a
# 40-line parser is easier to understand than pulling in a Markdown library.
# ---------------------------------------------------------------------------


@dataclass
class Block:
    kind: str  # "h1" | "h2" | "h3" | "para" | "bullet" | "table" | "pagebreak"
    text: str = ""
    rows: list[list[str]] = field(default_factory=list)  # only for "table"


def parse_markdown(source: str) -> list[Block]:
    blocks: list[Block] = []
    paragraph: list[str] = []  # lines of the paragraph being collected
    table: list[list[str]] = []  # rows of the table being collected

    def flush() -> None:
        # Close any open paragraph/table and append it as a block.
        if paragraph:
            blocks.append(Block("para", " ".join(paragraph)))
            paragraph.clear()
        if table:
            blocks.append(Block("table", rows=[row[:] for row in table]))
            table.clear()

    for raw_line in source.splitlines():
        line = raw_line.strip()
        if line == PAGE_BREAK:
            flush()
            blocks.append(Block("pagebreak"))
        elif not line:
            flush()
        elif line.startswith("|"):
            if paragraph:
                flush()
            cells = [cell.strip() for cell in line.strip("|").split("|")]
            # Skip the Markdown separator row: | --- | --- |
            if not all(set(cell) <= {"-", ":"} for cell in cells):
                table.append(cells)
        elif line.startswith("### "):
            flush()
            blocks.append(Block("h3", line[4:]))
        elif line.startswith("## "):
            flush()
            blocks.append(Block("h2", line[3:]))
        elif line.startswith("# "):
            flush()
            blocks.append(Block("h1", line[2:]))
        elif line.startswith("- "):
            flush()
            blocks.append(Block("bullet", line[2:]))
        else:
            paragraph.append(line)
    flush()
    return blocks


# ---------------------------------------------------------------------------
# PDF rendering with PyMuPDF.
#
# We lay text out ourselves (measure words, wrap lines, move a y-cursor down
# the page) instead of using an HTML-to-PDF helper. That keeps two properties
# the ingestion pipeline and the eval dataset rely on:
#   1. Pagination is exactly what the source says (one source page = one PDF
#      page), so an eval question can say "the answer is on page 3".
#   2. Font sizes are exactly what we choose. The PDF parser detects headings
#      by "this text is noticeably bigger than the body text", so headings
#      must never be shrunk to fit.
# ---------------------------------------------------------------------------

PAGE_WIDTH, PAGE_HEIGHT = 595, 842  # A4 in PDF points (1 pt = 1/72 inch)
MARGIN_X = 64
CONTENT_WIDTH = PAGE_WIDTH - 2 * MARGIN_X
BODY_TOP = 74  # body text starts below the running header
BODY_BOTTOM = 780  # body text must end above the footer
HEADER_BASELINE = 44
FOOTER_BASELINE = 812
BULLET_INDENT = 14


@dataclass(frozen=True)
class TextStyle:
    font: str  # PyMuPDF base-14 font code: "helv" = Helvetica, "hebo" = Helvetica-Bold
    size: float  # font size in points
    leading: float  # distance between consecutive baselines
    space_before: float  # extra gap above the block
    space_after: float  # extra gap below the block


PDF_STYLES = {
    "h1": TextStyle("hebo", 18, 23, 4, 6),
    "h2": TextStyle("hebo", 14, 18, 10, 3),
    "para": TextStyle("helv", 10.5, 14.2, 0, 7),
    "bullet": TextStyle("helv", 10.5, 14.2, 0, 2.5),
}
SMALL_FONT = ("helv", 8)  # running header + footer

# The base-14 PDF fonts only cover plain Latin characters; anything else (an
# em dash, a curly quote, a bullet symbol) renders as a dot. We translate the
# common offenders to ASCII instead of shipping font files.
ASCII_REPLACEMENTS = {"—": "-", "–": "-", "‘": "'", "’": "'",
                      "“": '"', "”": '"', "•": "-", " ": " "}


def to_pdf_safe(text: str) -> str:
    for bad, good in ASCII_REPLACEMENTS.items():
        text = text.replace(bad, good)
    if not text.isascii():
        raise ValueError(f"Non-ASCII character in PDF text: {text!r}")
    return text


def wrap_text(text: str, style: TextStyle, width: float) -> list[str]:
    """Greedy word wrap: keep adding words while the line still fits."""
    lines: list[str] = []
    current = ""
    for word in text.split():
        candidate = f"{current} {word}" if current else word
        if pymupdf.get_text_length(candidate, fontname=style.font, fontsize=style.size) <= width:
            current = candidate
        else:
            if current:
                lines.append(current)
            current = word
    if current:
        lines.append(current)
    return lines


def split_pages(blocks: list[Block]) -> list[list[Block]]:
    pages: list[list[Block]] = [[]]
    for block in blocks:
        if block.kind == "pagebreak":
            pages.append([])
        else:
            pages[-1].append(block)
    return [page for page in pages if page]


def render_pdf(spec: DocSpec, markdown: str, out_path: Path) -> int:
    pages = split_pages(parse_markdown(markdown))
    total = len(pages)
    pdf = pymupdf.open()

    for page_number, page_blocks in enumerate(pages, start=1):
        page = pdf.new_page(width=PAGE_WIDTH, height=PAGE_HEIGHT)

        # Running header FIRST, so it is the first line of the page's text.
        # (Text extraction follows drawing order, and real PDFs put the header
        # first, so the cleaner sees it where it would normally be.)
        if spec.running_header:
            page.insert_text((MARGIN_X, HEADER_BASELINE), to_pdf_safe(spec.running_header),
                             fontname=SMALL_FONT[0], fontsize=SMALL_FONT[1])
            page.draw_line((MARGIN_X, HEADER_BASELINE + 6),
                           (PAGE_WIDTH - MARGIN_X, HEADER_BASELINE + 6),
                           color=(0.6, 0.6, 0.6), width=0.5)

        y = BODY_TOP  # top of the remaining free space on this page
        previous_kind = None
        for block in page_blocks:
            if block.kind not in PDF_STYLES:
                raise ValueError(f"{spec.output}: '{block.kind}' blocks are not supported in PDFs")
            style = PDF_STYLES[block.kind]
            if previous_kind is not None:
                y += style.space_before
                if previous_kind == "bullet" and block.kind == "para":
                    y += 4  # small gap between a list and the paragraph after it

            x, width = MARGIN_X, CONTENT_WIDTH
            if block.kind == "bullet":
                x, width = MARGIN_X + BULLET_INDENT, CONTENT_WIDTH - BULLET_INDENT
                page.insert_text((MARGIN_X + 3, y + style.size), "-",
                                 fontname=style.font, fontsize=style.size)

            for line in wrap_text(to_pdf_safe(block.text), style, width):
                baseline = y + style.size
                if baseline > BODY_BOTTOM:
                    # Never shrink text to make it fit - fail loudly instead,
                    # so the source author moves content to the next page.
                    raise OverflowError(f"{spec.output}: page {page_number} overflows; "
                                        f"add a page break earlier or shorten it")
                page.insert_text((x, baseline), line, fontname=style.font, fontsize=style.size)
                y += style.leading
            y += style.space_after
            previous_kind = block.kind

        # Footer LAST, so "Page N of M" is the final line of the page's text.
        footer = f"Page {page_number} of {total}"
        footer_width = pymupdf.get_text_length(footer, fontname=SMALL_FONT[0], fontsize=SMALL_FONT[1])
        page.insert_text(((PAGE_WIDTH - footer_width) / 2, FOOTER_BASELINE), footer,
                         fontname=SMALL_FONT[0], fontsize=SMALL_FONT[1])

    pdf_date = FIXED_TIMESTAMP.strftime("D:%Y%m%d%H%M%S")
    pdf.set_metadata({
        "title": to_pdf_safe(spec.title),
        "author": "NEXA sample data generator",
        "subject": DISCLAIMER,
        "creator": "scripts/generate_sample_data.py",
        "producer": "PyMuPDF",
        "creationDate": pdf_date,
        "modDate": pdf_date,
    })
    # no_new_id keeps the PDF's internal /ID stable between runs (determinism).
    pdf.save(out_path, garbage=4, deflate=True, no_new_id=True)
    pdf.close()
    return total


# ---------------------------------------------------------------------------
# DOCX rendering with python-docx.
#
# We use Word's built-in "Heading 1/2/3" styles (doc.add_heading) rather than
# just making text bold. That is how real, well-made Word documents mark their
# structure, and it is what the DOCX parser reads to find section titles.
# ---------------------------------------------------------------------------


def render_docx(spec: DocSpec, markdown: str, out_path: Path) -> None:
    document = docx.Document()
    props = document.core_properties
    props.title = spec.title
    props.author = "NEXA sample data generator"
    props.comments = DISCLAIMER
    props.created = FIXED_TIMESTAMP
    props.modified = FIXED_TIMESTAMP
    props.last_modified_by = "NEXA sample data generator"
    props.revision = 1

    for block in parse_markdown(markdown):
        if block.kind in ("h1", "h2", "h3"):
            document.add_heading(block.text, level=int(block.kind[1]))
        elif block.kind == "para":
            document.add_paragraph(block.text)
        elif block.kind == "bullet":
            document.add_paragraph(block.text, style="List Bullet")
        elif block.kind == "table":
            header, *body = block.rows
            table = document.add_table(rows=1 + len(body), cols=len(header))
            table.style = "Table Grid"
            for col, text in enumerate(header):
                cell_run = table.rows[0].cells[col].paragraphs[0].add_run(text)
                cell_run.bold = True
            for row_index, row in enumerate(body, start=1):
                for col, text in enumerate(row):
                    table.rows[row_index].cells[col].text = text
        elif block.kind == "pagebreak":
            pass  # Word paginates on its own; DOCX files have no fixed pages.

    buffer = io.BytesIO()
    document.save(buffer)
    out_path.write_bytes(make_zip_deterministic(buffer.getvalue()))


def make_zip_deterministic(data: bytes) -> bytes:
    """A .docx is a ZIP archive, and ZIP entries carry a 'last modified' time.
    Rewrite every entry with a fixed time so the file hash never changes."""
    source = zipfile.ZipFile(io.BytesIO(data))
    output = io.BytesIO()
    with zipfile.ZipFile(output, "w", zipfile.ZIP_DEFLATED) as target:
        for info in source.infolist():  # keep the original entry order
            entry = zipfile.ZipInfo(info.filename, date_time=FIXED_TIMESTAMP.timetuple()[:6])
            entry.compress_type = zipfile.ZIP_DEFLATED
            entry.external_attr = 0o644 << 16
            target.writestr(entry, source.read(info.filename))
    return output.getvalue()


# ---------------------------------------------------------------------------
# The generated 50-page inspection report.
#
# Everything random comes from ONE random.Random(seed) instance, so the output
# is identical on every run. A few units are pinned with hand-written facts
# ("overrides") because other documents talk about them: Unit 4B and Unit 7A
# have leases, and one Building C unit has a dog, which the Building C house
# rules forbid - a nice cross-document fact for multi-document questions.
# ---------------------------------------------------------------------------

INSPECTION_SEED = 2026
INSPECTORS = ["Samuel Okafor", "Mei Tanaka", "Grace Holloway"]
FIRST_NAMES = ["Aaron", "Bianca", "Callum", "Deepa", "Elliot", "Fatima", "George", "Hana",
               "Isaac", "Jasmine", "Kofi", "Laura", "Mateo", "Nadia", "Oliver", "Phoebe",
               "Quentin", "Rosa", "Tariq", "Uma", "Victor", "Wen", "Yusuf", "Zara"]
LAST_NAMES = ["Abbott", "Brennan", "Castillo", "Dixon", "Evans", "Fischer", "Gupta", "Hughes",
              "Ivanova", "Jensen", "Kaur", "Lindqvist", "Mensah", "Novak", "O'Brien", "Petrov",
              "Quinn", "Rossi", "Sato", "Turner", "Vega", "Walsh", "Yilmaz", "Zhou"]
ROOMS = ["Kitchen", "Bathroom", "Living area", "Bedrooms", "Laundry", "Balcony"]

# (room, issue description, contractor or None, action). A None contractor
# means the issue was fixed on the spot or is the tenant's job.
ISSUE_POOL = [
    ("Kitchen", "Dripping kitchen mixer tap.", "Coastline Plumbing", "replace the tap cartridge"),
    ("Kitchen", "Rangehood light not working.", "Brightwire Electrical", "replace the rangehood globe and switch"),
    ("Kitchen", "Evidence of cockroaches in the cupboard under the sink.", "Evergreen Pest Control", "carry out a pest treatment"),
    ("Bathroom", "Mould on the bathroom ceiling above the shower.", "FreshAir Mould Solutions", "treat the mould and check the exhaust fan"),
    ("Bathroom", "Bathroom exhaust fan not working.", "Brightwire Electrical", "replace the exhaust fan motor"),
    ("Bathroom", "Shower screen seal is cracked and leaking onto the floor.", "Coastline Plumbing", "reseal the shower screen"),
    ("Living area", "Split-system air conditioner not cooling.", "Brightwire Electrical", "service the air conditioner"),
    ("Living area", "Carpet stain of about 30 cm near the balcony door.", None, "Tenant advised to arrange carpet cleaning before the next inspection."),
    ("Bedrooms", "Water stain on the main bedroom ceiling.", "Coastline Plumbing", "investigate a possible leak from the unit above"),
    ("Bedrooms", "Wardrobe sliding door has come off its track.", "Harbour Handyman Services", "refit the wardrobe door"),
    ("Laundry", "Cracked floor tile at the laundry entrance.", "Harbour Handyman Services", "replace the cracked tile"),
    ("Laundry", "Washing machine tap leaking at the connection.", "Coastline Plumbing", "replace the washing machine tap washer"),
    ("Balcony", "Balcony sliding door lock is stiff and hard to engage.", "SecureLock Locksmiths", "service the sliding door lock"),
    ("Balcony", "Blocked balcony floor drain with pooling water.", "Coastline Plumbing", "clear the balcony drain"),
]
SMOKE_ALARM_FAULT = "Hallway smoke alarm failed the test."
SMOKE_ALARM_ACTION = "Smoke alarm battery replaced on site by the inspector; alarm retested and working."
NOTES_POOL = [  # usable whether or not the tenant was home
    "The tenant keeps the unit clean and tidy.",
    "The unit is presented in good order overall.",
    "Minor wear and tear consistent with the age of the fittings.",
    "Windows and window tracks are clean. No signs of damp apart from the items noted.",
]
NOTES_TENANT_PRESENT = [  # only make sense if the tenant attended
    "The tenant asked about repainting the second bedroom; referred to the property manager.",
    "The tenant reported no other concerns during the inspection.",
]

# Units whose details are pinned because other documents mention them.
INSPECTION_OVERRIDES = {
    ("A", "4B"): dict(
        tenant="Jordan Ellis", date=dt.date(2026, 6, 18), inspector="Samuel Okafor",
        tenant_present=True, smoke_fault=False,
        issues=[("Bathroom", "Mould on the bathroom ceiling above the shower.",
                 "FreshAir Mould Solutions", "treat the mould and check the exhaust fan")],
        note="An indoor cat is kept at the unit as approved under Clause 10.1 of the lease. No pet damage observed.",
    ),
    ("C", "7A"): dict(
        tenant="Priya Raman", date=dt.date(2026, 4, 9), inspector="Mei Tanaka",
        tenant_present=False, smoke_fault=True,
        issues=[("Balcony", "Blocked balcony floor drain with pooling water.",
                 "Coastline Plumbing", "clear the balcony drain")],
        note="The tenant was not home; entry was made with the agent's key after 7 days written notice.",
    ),
    ("C", "3D"): dict(
        tenant="Callum Brennan", date=dt.date(2026, 5, 20), inspector="Grace Holloway",
        tenant_present=True, smoke_fault=False,
        issues=[("Living area", "A dog was observed living in the unit. Dogs are not permitted under the Building C House Rules.",
                 None, "Breach notice issued to the tenant under the Building C House Rules; the dog must be removed within 14 days.")],
        note="The tenant said the dog belongs to a relative and is staying temporarily.",
    ),
}


def month_after(date: dt.date, months: int) -> str:
    month_index = date.month - 1 + months
    year = date.year + month_index // 12
    return dt.date(year, month_index % 12 + 1, 1).strftime("%B %Y")


def format_date(date: dt.date) -> str:
    return f"{date.day} {date.strftime('%B %Y')}"  # e.g. "18 June 2026"


def build_inspection_reports_source() -> str:
    rng = random.Random(INSPECTION_SEED)

    # 1. Choose 50 distinct units: 18 in Building A, 16 in B, 16 in C.
    #    A unit id is floor number + door letter, e.g. "4B" = level 4, door B.
    #    Pinned unit ids (4B, 7A, 3D) are kept out of the OTHER buildings, so
    #    "Unit 7A" always means the Building C apartment from the lease.
    units: list[tuple[str, str]] = []
    all_pinned = {unit for (_b, unit) in INSPECTION_OVERRIDES}
    for building, count in (("A", 18), ("B", 16), ("C", 16)):
        candidates = [f"{floor}{door}" for floor in range(1, 10) for door in "ABCD"]
        pinned = [unit for (b, unit) in INSPECTION_OVERRIDES if b == building]
        others = [u for u in candidates if u not in all_pinned]
        chosen = pinned + rng.sample(others, count - len(pinned))
        # Sort by floor then door so the report reads like a building walk-through.
        chosen.sort(key=lambda unit: (int(unit[:-1]), unit[-1]))
        units += [(building, unit) for unit in chosen]

    # 2. Unique inspection references and work order numbers. They are shuffled
    #    so page order does not reveal the number (exact-identifier questions
    #    should need real retrieval, not arithmetic).
    references = rng.sample(range(101, 400), len(units))
    work_orders = iter(rng.sample(range(6001, 6999), 200))

    # 3. Unique tenant names (pinned names are reserved first).
    reserved = {o["tenant"] for o in INSPECTION_OVERRIDES.values()}
    all_names = [f"{f} {l}" for f in FIRST_NAMES for l in LAST_NAMES]
    tenant_names = iter(n for n in rng.sample(all_names, 80) if n not in reserved)

    start = dt.date(2026, 1, 12)
    pages: list[str] = []
    for index, (building, unit) in enumerate(units):
        override = INSPECTION_OVERRIDES.get((building, unit))
        if override:
            tenant, date, inspector = override["tenant"], override["date"], override["inspector"]
            tenant_present, smoke_fault = override["tenant_present"], override["smoke_fault"]
            issues, note = override["issues"], override["note"]
        else:
            tenant = next(tenant_names)
            date = start + dt.timedelta(days=rng.randrange(0, 225))
            while date.weekday() >= 5:  # inspections happen on weekdays
                date += dt.timedelta(days=1)
            inspector = rng.choice(INSPECTORS)
            tenant_present = rng.random() < 0.7
            smoke_fault = rng.random() < 0.15
            issue_count = rng.choice([0, 0, 1, 1, 1, 2, 2, 3])
            issues = rng.sample(ISSUE_POOL, issue_count)
            note = rng.choice(NOTES_POOL + (NOTES_TENANT_PRESENT if tenant_present else []))

        # Room ratings: rooms with an issue drop to Fair or Poor.
        ratings = {room: "Good" for room in ROOMS}
        for room, *_ in issues:
            ratings[room] = rng.choice(["Fair", "Poor"]) if not override else "Fair"
        problem_count = len(issues) + (1 if smoke_fault else 0)
        overall = ["Good", "Good", "Fair", "Needs attention"][min(problem_count, 3)]

        # Follow-up actions: a work order per contractor issue, due 14 days
        # later (matches the "routine repairs within 14 days" maintenance rule).
        actions: list[str] = []
        for _room, _issue, contractor, action in issues:
            if contractor:
                due = date + dt.timedelta(days=14)
                actions.append(f"Work order WO-{next(work_orders)} raised with {contractor} to {action}, "
                               f"due by {format_date(due)}.")
            else:
                actions.append(action)
        if smoke_fault:
            actions.append(SMOKE_ALARM_ACTION)
        actions.append(f"Next routine inspection due: {month_after(date, 6)}.")

        issue_lines = [f"- {room}: {issue}" for room, issue, *_ in issues]
        if smoke_fault:
            issue_lines.append(f"- Smoke alarms: {SMOKE_ALARM_FAULT}")
        if not issue_lines:
            issue_lines = ["- No issues found. The unit is well maintained."]

        smoke_status = "One alarm failed the test (see Issues Found)" if smoke_fault else "Tested, all working"
        reference = f"INS-2026-{references[index]:04d}"
        lines = [f"# Unit {unit}, Building {building}"]
        if index == 0:
            lines += ["", DISCLAIMER + " Routine inspection reports for Harbourview Residences, "
                      "January to August 2026. One report per page."]
        lines += [
            "",
            f"Routine inspection report {reference} for Unit {unit}, Building {building}, "
            f"Harbourview Residences, 42 Quay Street, Port Ellison.",
            "",
            "## Inspection Details",
            f"- Inspection reference: {reference}",
            f"- Date of inspection: {format_date(date)}",
            f"- Inspector: {inspector}",
            f"- Tenant: {tenant}",
            f"- Tenant present: {'Yes' if tenant_present else 'No'}",
            f"- Overall condition: {overall}",
            "",
            "## Condition by Room",
            *[f"- {room}: {rating}" for room, rating in ratings.items()],
            f"- Smoke alarms: {smoke_status}",
            "",
            "## Issues Found",
            *issue_lines,
            "",
            "## Follow-up Actions",
            *[f"- {action}" for action in actions],
            "",
            "## Inspector Notes",
            note,
        ]
        pages.append("\n".join(lines))
    return f"\n\n{PAGE_BREAK}\n\n".join(pages) + "\n"


# ---------------------------------------------------------------------------
# Main: render every document and print a short summary.
# ---------------------------------------------------------------------------


def describe_pdf(path: Path) -> str:
    """Summarise a PDF's font sizes so you can see headings are bigger than
    body text (which is what the PDF parser's heading detection relies on)."""
    sizes: dict[float, int] = {}
    with pymupdf.open(path) as pdf:
        for page in pdf:
            for block in page.get_text("dict")["blocks"]:
                for line in block.get("lines", []):
                    for span in line["spans"]:
                        size = round(span["size"], 1)
                        sizes[size] = sizes.get(size, 0) + len(span["text"])
    return ", ".join(f"{size}pt x{chars}" for size, chars in sorted(sizes.items(), reverse=True))


def main() -> None:
    for spec in DOCUMENTS:
        out_dir = OUTPUT_DIR / spec.domain
        out_dir.mkdir(parents=True, exist_ok=True)
        out_path = out_dir / spec.output

        if spec.source is None:
            text = build_inspection_reports_source()
        else:
            text = (SOURCE_DIR / spec.domain / spec.source).read_text(encoding="utf-8")

        suffix = out_path.suffix
        if suffix == ".pdf":
            pages = render_pdf(spec, text, out_path)
            detail = f"{pages} pages; font sizes: {describe_pdf(out_path)}"
        elif suffix == ".docx":
            render_docx(spec, text, out_path)
            detail = "Word document with Heading styles"
        elif suffix in (".md", ".txt"):
            # Markdown and plain-text documents are shipped as written.
            out_path.write_text(text.rstrip() + "\n", encoding="utf-8")
            detail = "copied as-is"
        else:
            raise ValueError(f"Unsupported output format: {spec.output}")

        digest = hashlib.sha256(out_path.read_bytes()).hexdigest()[:12]
        print(f"{spec.domain}/{spec.output:<36} sha256:{digest}  {detail}")


if __name__ == "__main__":
    main()
